# SPDX-License-Identifier: Apache-2.0
"""Render Markdown fixture sources to PDF, DOCX, HTML, MD and TXT with expected elements.

Expected elements derive from the sources (and the PDF page log recorded while
reportlab draws), never from footnote's parsers.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import (
    Flowable,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

HERE = Path(__file__).resolve().parent
SOURCES_DIR = HERE / "sources"
DEFAULT_OUT = HERE / "generated"
EXTENSIONS = ("pdf", "docx", "html", "md", "txt")
HEADING_SIZES = {1: 20.0, 2: 16.0, 3: 13.0, 4: 12.0, 5: 11.5, 6: 11.2}
BODY_SIZE = 11.0
FIXED_TIME = dt.datetime(2026, 1, 1, tzinfo=dt.UTC)
FRONT_MATTER = "---\ntitle: Ignored front matter\ntags: [fixture]\n---\n\n"


@dataclass(frozen=True)
class Block:
    """One block of the restricted Markdown subset."""

    kind: str  # heading | paragraph | list_item | table | code
    text: str
    level: int | None = None
    rows: tuple[tuple[str, ...], ...] = field(default=())


def _split_row(line: str) -> tuple[str, ...]:
    return tuple(c.strip() for c in line.strip().strip("|").split("|"))


def _is_separator(cells: tuple[str, ...]) -> bool:
    return all(c and set(c) <= {"-", ":"} for c in cells)


def table_text(rows: tuple[tuple[str, ...], ...]) -> str:
    """Markdown table text: header, `|---|` separator, body rows."""
    head = "| " + " | ".join(rows[0]) + " |"
    sep = "|" + "|".join("---" for _ in rows[0]) + "|"
    body = ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join([head, sep, *body])


def parse_source(text: str) -> list[Block]:
    """Parse the restricted Markdown subset into blocks."""
    blocks: list[Block] = []
    para: list[str] = []
    lines = text.splitlines()
    i = 0

    def flush() -> None:
        if para:
            blocks.append(Block("paragraph", " ".join(para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(code)))
        elif stripped.startswith("#"):
            flush()
            hashes = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(Block("heading", stripped[hashes:].strip(), level=hashes))
        elif stripped.startswith("- "):
            flush()
            blocks.append(Block("list_item", stripped[2:].strip()))
        elif stripped.startswith("|"):
            flush()
            rows: list[tuple[str, ...]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_row(lines[i])
                if not _is_separator(cells):
                    rows.append(cells)
                i += 1
            i -= 1
            tup = tuple(rows)
            blocks.append(Block("table", table_text(tup), rows=tup))
        elif not stripped:
            flush()
        else:
            para.append(stripped)
        i += 1
    flush()
    return blocks


def _element(block: Block, order: int, path: tuple[str, ...], **extra: Any) -> dict[str, Any]:
    return {
        "kind": block.kind,
        "text": block.text,
        "level": block.level,
        "heading_path": list(path),
        "page": extra.get("page"),
        "page_to": extra.get("page_to"),
        "order": order,
    }


def expected_elements(
    blocks: list[Block], pages: dict[int, tuple[int, int]] | None = None
) -> list[dict[str, Any]]:
    """Expected elements; `pages` maps block index to (page, page_to) for PDFs."""
    out: list[dict[str, Any]] = []
    stack: list[tuple[int, str]] = []
    for idx, block in enumerate(blocks):
        if block.kind == "heading":
            assert block.level is not None
            while stack and stack[-1][0] >= block.level:
                stack.pop()
            stack.append((block.level, block.text))
        path = tuple(t for _, t in stack)
        page, page_to = pages[idx] if pages else (None, None)
        out.append(_element(block, len(out), path, page=page, page_to=page_to))
    return out


def txt_blocks(blocks: list[Block]) -> list[Block]:
    """TXT keeps only prose: headings and list items become plain paragraphs."""
    return [
        Block("paragraph", b.text)
        for b in blocks
        if b.kind in {"heading", "paragraph", "list_item"}
    ]


# ---------------------------------------------------------------- PDF


class _Doc(SimpleDocTemplate):
    """Doc template logging the pages every block is drawn on."""

    def __init__(self, filename: str, **kw: Any) -> None:
        super().__init__(filename, **kw)
        self.page_log: dict[int, tuple[int, int]] = {}

    def afterFlowable(self, flowable: Flowable) -> None:  # noqa: N802 (reportlab hook)
        idx = getattr(flowable, "fn_index", None)
        if idx is None:
            style = getattr(flowable, "style", None)
            name = getattr(style, "name", "")
            if not name.startswith("fn"):
                return
            idx = int(name[2:])
        page = self.canv.getPageNumber()
        first, last = self.page_log.get(idx, (page, page))
        self.page_log[idx] = (min(first, page), max(last, page))


class _FnTable(Table):
    fn_index: int | None = None


def _escape(text: str) -> str:
    return html.escape(text, quote=False)


def render_pdf(blocks: list[Block], path: Path) -> dict[int, tuple[int, int]]:
    """Render PDF; return the page log (block index -> (page, page_to))."""
    doc = _Doc(
        str(path),
        pagesize=A4,
        invariant=1,
        pageCompression=0,
        title=path.stem,
        author="footnote fixtures",
        leftMargin=56,
        rightMargin=56,
        topMargin=56,
        bottomMargin=56,
    )
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=BODY_SIZE, leading=15)
    story: list[Flowable] = []
    for idx, block in enumerate(blocks):
        name = f"fn{idx}"
        if block.kind == "heading":
            assert block.level is not None
            size = HEADING_SIZES[block.level]
            style = ParagraphStyle(
                name,
                fontName="Helvetica-Bold",
                fontSize=size,
                leading=size + 4,
                spaceBefore=10,
                spaceAfter=6,
            )
            story.append(Paragraph(_escape(block.text), style))
        elif block.kind == "paragraph":
            style = ParagraphStyle(name, parent=body, spaceAfter=8)
            story.append(Paragraph(_escape(block.text), style))
        elif block.kind == "list_item":
            style = ParagraphStyle(name, parent=body, leftIndent=18, spaceAfter=3)
            story.append(Paragraph("• " + _escape(block.text), style))
        elif block.kind == "code":
            style = ParagraphStyle(name, fontName="Courier", fontSize=10, leading=12)
            story.append(Preformatted(block.text, style))
            story.append(Spacer(1, 8))
        else:
            table = _FnTable([list(r) for r in block.rows])
            table.fn_index = idx
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.8, (0, 0, 0)),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), BODY_SIZE),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 10))
    doc.build(story)
    return doc.page_log


# ---------------------------------------------------------------- DOCX


def render_docx(blocks: list[Block], path: Path) -> None:
    """Render DOCX with Heading N / List Bullet styles, grid tables and Code style."""
    doc = Document()
    code_style = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Courier New"
    code_style.font.size = Pt(10)
    for block in blocks:
        if block.kind == "heading":
            assert block.level is not None
            doc.add_paragraph(block.text, style=f"Heading {block.level}")
        elif block.kind == "paragraph":
            doc.add_paragraph(block.text)
        elif block.kind == "list_item":
            doc.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "code":
            for line in block.text.split("\n"):
                doc.add_paragraph(line, style="Code")
        else:
            table = doc.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(block.rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    props = doc.core_properties
    props.author = "footnote fixtures"
    props.last_modified_by = "footnote fixtures"
    props.title = path.stem
    props.created = FIXED_TIME.replace(tzinfo=None)
    props.modified = FIXED_TIME.replace(tzinfo=None)
    props.revision = 1
    doc.save(str(path))


# ---------------------------------------------------------------- HTML / MD / TXT


def render_html(blocks: list[Block], title: str) -> str:
    """HTML with nav, footer and script boilerplate that parsers must strip."""
    parts = [
        "<!doctype html>",
        '<html lang="en"><head><meta charset="utf-8">',
        f"<title>{html.escape(title)}</title>",
        "<script>var tracking = 'strip me';</script></head><body>",
        '<nav><a href="/">Home navigation</a> <a href="/about">About navigation</a></nav>',
        "<main>",
    ]
    in_list = False
    for block in blocks:
        if in_list and block.kind != "list_item":
            parts.append("</ul>")
            in_list = False
        if block.kind == "heading":
            parts.append(f"<h{block.level}>{html.escape(block.text)}</h{block.level}>")
        elif block.kind == "paragraph":
            parts.append(f"<p>{html.escape(block.text)}</p>")
        elif block.kind == "list_item":
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{html.escape(block.text)}</li>")
        elif block.kind == "code":
            parts.append(f"<pre><code>{html.escape(block.text)}</code></pre>")
        else:
            rows = []
            for r, row in enumerate(block.rows):
                tag = "th" if r == 0 else "td"
                cells = "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row)
                rows.append(f"<tr>{cells}</tr>")
            parts.append("<table>" + "".join(rows) + "</table>")
    if in_list:
        parts.append("</ul>")
    parts += [
        "</main>",
        "<footer>Footer boilerplate to strip</footer>",
        "<script>console.log('strip me too');</script>",
        "</body></html>",
    ]
    return "\n".join(parts) + "\n"


def render_txt(blocks: list[Block]) -> str:
    """Paragraphs separated by blank lines."""
    return "\n\n".join(b.text for b in txt_blocks(blocks)) + "\n"


# ---------------------------------------------------------------- driver


def _write_json(path: Path, elements: list[dict[str, Any]]) -> None:
    path.write_text(
        json.dumps(elements, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def generate(out_dir: Path) -> list[Path]:
    """Write every fixture and its `.expected.json` into `out_dir`; return the paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    def emit(path: Path, elements: list[dict[str, Any]]) -> None:
        expected = path.with_name(path.name + ".expected.json")
        _write_json(expected, elements)
        written.extend([path, expected])

    for src in sorted(SOURCES_DIR.glob("*.md")):
        text = src.read_text(encoding="utf-8")
        blocks = parse_source(text)
        name = src.stem
        plain = expected_elements(blocks)

        pdf = out_dir / f"{name}.pdf"
        log = render_pdf(blocks, pdf)
        emit(pdf, expected_elements(blocks, log))

        docx = out_dir / f"{name}.docx"
        render_docx(blocks, docx)
        emit(docx, plain)

        page = out_dir / f"{name}.html"
        page.write_text(render_html(blocks, name), encoding="utf-8")
        emit(page, plain)

        md = out_dir / f"{name}.md"
        md.write_text(FRONT_MATTER + text, encoding="utf-8")
        emit(md, plain)

        txt = out_dir / f"{name}.txt"
        txt.write_text(render_txt(blocks), encoding="utf-8")
        emit(txt, expected_elements(txt_blocks(blocks)))
    return written


def main() -> None:
    """CLI entry point."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    for path in generate(args.out):
        print(path)


if __name__ == "__main__":
    main()
